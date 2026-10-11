"""M5 — critères d'acceptation : cycle de vie, numérotation, rendu, correctifs."""
from __future__ import annotations

from datetime import date
from decimal import Decimal

import openpyxl
import pytest

from core.entries_service import creer_entree
from core.enums import ModeNumerotation, StatutFacture, TypeActivite, Unite
from core.invoice_export import (
    CAPACITE_LABOR,
    CAPACITE_ODC,
    LABOR_PREMIERE,
    MODELE,
    generer_facture,
    nom_fichier_facture,
    totaux_attendus,
)
from core.invoice_service import (
    ajouter_odc,
    annuler_brouillon,
    attacher_entrees,
    code_facture_defaut,
    creer_brouillon,
    detacher_entrees,
    emettre,
    entrees_facturables,
    marquer_payee,
    prochain_numero,
    supprimer_odc,
)
from core.models import Entry, Invoice, InvoiceLabor, InvoiceOdc
from core.xlsx_utils import TEMPLATES_DIR, charger_modele


def ouvrir(db, invoice_id):
    with db.get_session() as session:
        flux, nom = generer_facture(session, invoice_id)
    return openpyxl.load_workbook(flux), nom


# --------------------------------------------------------------------------- #
# Numérotation
# --------------------------------------------------------------------------- #
def test_code_facture_par_defaut():
    assert code_facture_defaut("Jane Doe") == "JD"
    assert code_facture_defaut("Zakaria Ben Hamouda") == "ZBH"
    assert code_facture_defaut("") == "XX"


def test_numero_par_expert_et_par_annee(db, jeu, facture_emise):
    with db.get_session() as session:
        assert session.get(Invoice, facture_emise).numero == "TC-JD-2026-001"
        # La séquence suivante du même expert, même année.
        assert prochain_numero(session, jeu["expert"], 2026) == "TC-JD-2026-002"
        # Un autre expert repart à 001, sous son propre code.
        assert prochain_numero(session, jeu["autre"], 2026) == "TC-JR-2026-001"
        # Année suivante : la séquence repart.
        assert prochain_numero(session, jeu["expert"], 2027) == "TC-JD-2026-002".replace(
            "2026-002", "2027-001"
        )


def test_numeros_de_deux_experts_restent_uniques(db, jeu, entrees, facture_emise):
    """Deux experts facturant la même année n'entrent pas en collision."""
    with db.get_session() as session:
        _, erreurs = creer_entree(
            session, jeu["autre"], jeu["autre"], jeu["projet"], date(2026, 10, 8),
            Unite.HEURE, Decimal("5"), TypeActivite.FEDERAL,
        )
        assert not erreurs
        ids = [e.id for e in entrees_facturables(session, jeu["autre"])]
        facture, erreurs = creer_brouillon(
            session, jeu["autre"], jeu["autre"], entry_ids=ids
        )
        assert not erreurs
        autre_id = facture.id
    with db.get_session() as session:
        facture, erreurs = emettre(session, jeu["autre"], autre_id, date(2026, 10, 31))
        assert not erreurs, erreurs
        assert facture.numero == "TC-JR-2026-001"
    with db.get_session() as session:
        numeros = [i.numero for i in session.query(Invoice).all()]
        assert len(numeros) == len(set(numeros))


def test_mode_de_numerotation_globale(db, jeu, entrees):
    from core.db import get_config

    with db.get_session() as session:
        get_config(session).mode_numerotation = ModeNumerotation.GLOBALE
    with db.get_session() as session:
        assert prochain_numero(session, jeu["expert"], 2026) == "TC-2026-001"


# --------------------------------------------------------------------------- #
# Cycle de vie
# --------------------------------------------------------------------------- #
def test_brouillon_reserve_les_entrees(db, jeu, brouillon, entrees):
    with db.get_session() as session:
        assert all(session.get(Entry, eid).invoice_id == brouillon for eid in entrees)
        # Elles ne sont plus proposées pour une autre facture.
        assert entrees_facturables(session, jeu["expert"]) == []


def test_annulation_du_brouillon_libere_les_entrees(db, jeu, brouillon, entrees):
    with db.get_session() as session:
        erreurs = annuler_brouillon(session, jeu["expert"], brouillon)
        assert not erreurs
    with db.get_session() as session:
        assert session.get(Invoice, brouillon) is None
        assert all(session.get(Entry, eid).invoice_id is None for eid in entrees)
        assert len(entrees_facturables(session, jeu["expert"])) == 3


def test_detachement_recalcule_le_brouillon(db, jeu, brouillon, entrees):
    with db.get_session() as session:
        erreurs = detacher_entrees(session, jeu["expert"], brouillon, [entrees[1]])
        assert not erreurs
    with db.get_session() as session:
        facture = session.get(Invoice, brouillon)
        assert facture.sous_total_labor == Decimal("1500.00")  # 2 x 750
        assert facture.total_net == Decimal("1250.00")
        assert session.get(Entry, entrees[1]).invoice_id is None


def test_emission_fige_les_lignes_et_verrouille(db, jeu, facture_emise, entrees):
    with db.get_session() as session:
        facture = session.get(Invoice, facture_emise)
        assert StatutFacture(facture.statut) is StatutFacture.EMISE
        assert facture.numero == "TC-JD-2026-001"
        assert facture.date_emission == date(2026, 10, 31)
        assert facture.date_echeance == date(2026, 11, 30)
        assert len(facture.labor_lines) == 3
        assert facture.sous_total_labor == Decimal("3100.00")
        assert facture.total_net == Decimal("2850.00")
        assert [l.item for l in facture.labor_lines] == [1, 2, 3]
        # Une ligne LABOR par entrée, sans regroupement.
        assert sorted(float(l.quantite) for l in facture.labor_lines) == [2.0, 7.5, 7.5]
        assert all(session.get(Entry, eid).invoice_id == facture_emise for eid in entrees)


def test_facture_emise_non_reemissible(db, jeu, facture_emise):
    with db.get_session() as session:
        _, erreurs = emettre(session, jeu["expert"], facture_emise)
        assert erreurs and "déjà" in erreurs[0]
        erreurs = annuler_brouillon(session, jeu["expert"], facture_emise)
        assert erreurs
        _, erreurs = ajouter_odc(
            session, jeu["expert"], facture_emise, "P", "x", 1, Decimal("1")
        )
        assert erreurs


def test_facture_vide_refusee(db, jeu):
    with db.get_session() as session:
        facture, erreurs = creer_brouillon(session, jeu["expert"], jeu["expert"])
        assert not erreurs
        invoice_id = facture.id
    with db.get_session() as session:
        _, erreurs = emettre(session, jeu["expert"], invoice_id)
        assert erreurs and "vide" in erreurs[0]


def test_avance_superieure_au_total_refusee(db, jeu, brouillon):
    from core.invoice_service import definir_avance

    with db.get_session() as session:
        erreurs = definir_avance(session, jeu["expert"], brouillon, Decimal("99999"))
        assert not erreurs  # acceptée au brouillon
    with db.get_session() as session:
        _, erreurs = emettre(session, jeu["expert"], brouillon)
        assert erreurs and "dépasse le total" in erreurs[0]


def test_entree_non_facturable_refusee(db, jeu, entrees, brouillon):
    from core.entries_service import modifier_entree

    with db.get_session() as session:
        detacher_entrees(session, jeu["expert"], brouillon, [entrees[0]])
        modifier_entree(session, jeu["expert"], entrees[0], facturable=False)
    with db.get_session() as session:
        erreurs = attacher_entrees(session, jeu["expert"], brouillon, [entrees[0]])
        assert erreurs and "non facturable" in erreurs[0]


def test_entree_d_un_autre_expert_refusee(db, jeu, brouillon):
    with db.get_session() as session:
        autre, _ = creer_entree(
            session, jeu["autre"], jeu["autre"], jeu["projet"], date(2026, 10, 7),
            Unite.HEURE, Decimal("2"), TypeActivite.FEDERAL,
        )
        entry_id = autre.id
    with db.get_session() as session:
        erreurs = attacher_entrees(session, jeu["expert"], brouillon, [entry_id])
        assert erreurs and "autre expert" in erreurs[0]


def test_marquage_paye(db, jeu, facture_emise):
    with db.get_session() as session:
        facture, erreurs = marquer_payee(session, jeu["expert"], facture_emise, date(2026, 12, 1))
        assert not erreurs
        assert StatutFacture(facture.statut) is StatutFacture.PAYEE
        assert "2026-12-01" in facture.note
    with db.get_session() as session:
        _, erreurs = marquer_payee(session, jeu["expert"], facture_emise)
        assert erreurs


# --------------------------------------------------------------------------- #
# ODC
# --------------------------------------------------------------------------- #
def test_odc_montant_et_items(db, jeu, brouillon):
    with db.get_session() as session:
        for index in range(3):
            _, erreurs = ajouter_odc(
                session, jeu["expert"], brouillon, "Education Reform Study",
                f"Airfare leg {index + 1}", Decimal("2"), Decimal("431.75"),
            )
            assert not erreurs, erreurs
    with db.get_session() as session:
        facture = session.get(Invoice, brouillon)
        lignes = sorted(facture.odc_lines, key=lambda o: o.item)
        assert [l.item for l in lignes] == [1, 2, 3]
        assert all(l.montant == Decimal("863.50") for l in lignes)
        assert facture.sous_total_odc == Decimal("2590.50")
        assert facture.total_net == Decimal("3100.00") + Decimal("2590.50") - Decimal("250")


def test_suppression_odc_renumerote(db, jeu, brouillon):
    with db.get_session() as session:
        for index in range(3):
            ajouter_odc(
                session, jeu["expert"], brouillon, "P", f"odc{index}", 1, Decimal("10")
            )
    with db.get_session() as session:
        facture = session.get(Invoice, brouillon)
        cible = sorted(facture.odc_lines, key=lambda o: o.item)[1].id
        erreurs = supprimer_odc(session, jeu["expert"], cible)
        assert not erreurs
    with db.get_session() as session:
        facture = session.get(Invoice, brouillon)
        assert [l.item for l in sorted(facture.odc_lines, key=lambda o: o.item)] == [1, 2]


def test_odc_valide_ses_entrees(db, jeu, brouillon):
    with db.get_session() as session:
        _, erreurs = ajouter_odc(session, jeu["expert"], brouillon, "", "", 0, Decimal("-1"))
        assert len(erreurs) == 4


# --------------------------------------------------------------------------- #
# Rendu Excel
# --------------------------------------------------------------------------- #
def test_nom_de_fichier():
    assert nom_fichier_facture("TC-JD-2026-001", 1) == "Invoice_TC-JD-2026-001.xlsx"
    assert nom_fichier_facture(None, 7) == "Invoice_brouillon-7.xlsx"


def test_entetes_de_la_facture(db, jeu, facture_emise):
    ws = ouvrir(db, facture_emise)[0].active
    assert ws["B2"].value == "CONSULTANT INVOICE"
    assert ws["G2"].value == "Doe Consulting LLC"
    assert ws["G3"].value == "12 Elm Street"
    assert ws["G4"].value == "Raleigh, NC 27601"
    assert ws["G5"].value == "+1 919 555 0101"
    assert ws["G6"].value == "jane@doeconsulting.com"
    assert ws["G9"].value == "Bill To"
    assert ws["G10"].value == "Tangerine Central, LLC"
    assert ws["G11"].value == "Durham, NC, USA"
    assert ws["G12"].value == "carmen@tangerinecentral.org"
    assert ws["C10"].value == "TC-JD-2026-001"
    # openpyxl relit une cellule de date en datetime.
    assert ws["C11"].value.date() == date(2026, 10, 31)
    assert ws["C12"].value.date() == date(2026, 11, 30)
    assert ws["C11"].number_format == "mm-dd-yy"


def test_lignes_labor_et_formules(db, jeu, facture_emise):
    ws = ouvrir(db, facture_emise)[0].active
    assert ws.cell(row=LABOR_PREMIERE, column=2).value == 1
    assert ws.cell(row=LABOR_PREMIERE, column=3).value == "Education Reform Study"
    # Le livrable prime sur la note quand il existe.
    assert ws.cell(row=LABOR_PREMIERE, column=4).value == "Inception report"
    assert ws.cell(row=LABOR_PREMIERE, column=5).value == 7.5
    assert ws.cell(row=LABOR_PREMIERE, column=6).value == 100.0
    # Amount reste la formule du modèle, à référence structurée.
    assert "InvoiceDetails" in ws.cell(row=LABOR_PREMIERE, column=7).value
    # Les sous-totaux du modèle sont conservés.
    assert ws["E24"].value == "=SUBTOTAL(109,InvoiceDetails[Hours])"
    assert ws["G24"].value == '=IFERROR(SUM(InvoiceDetails[Amount]),"0")'


def test_quantite_decimale_lisible(db, jeu, facture_emise):
    """Le format « 0 » du modèle afficherait 8 pour une quantité de 7,5."""
    ws = ouvrir(db, facture_emise)[0].active
    assert charger_modele(MODELE).active["E16"].number_format == "0"
    assert ws.cell(row=LABOR_PREMIERE, column=5).number_format == "#,##0.##"


def test_correctif_amount_odc(db, jeu, brouillon):
    """Le modèle référence le tableau LABOR dans G27:G29 : corrigé en E*F."""
    with db.get_session() as session:
        ajouter_odc(
            session, jeu["expert"], brouillon, "Education Reform Study", "Airfare",
            Decimal("3"), Decimal("50"),
        )
    modele = charger_modele(MODELE).active
    assert "InvoiceDetails" in modele["G27"].value  # bug d'origine
    ws = ouvrir(db, brouillon)[0].active
    assert ws["G27"].value == "=E27*F27"


def test_correctif_sous_total_odc(db, jeu, brouillon):
    """Le modèle somme la colonne Amount du tableau LABOR en G30."""
    modele = charger_modele(MODELE).active
    assert modele["G30"].value == '=IFERROR(SUM(InvoiceDetails[Amount]),"0")'
    assert modele["G30"].value == modele["G24"].value  # d'où le double comptage
    ws = ouvrir(db, brouillon)[0].active
    assert ws["G30"].value == "=SUM(G27:G29)"
    assert ws["G31"].value == "=G24+G30"
    assert ws["G33"].value == "=G31-G32"


def test_totaux_attendus_coherents(db, jeu, facture_emise):
    with db.get_session() as session:
        attendus = totaux_attendus(session, facture_emise)
        facture = session.get(Invoice, facture_emise)
    assert attendus["sous_total_labor"] == Decimal("3100.00")
    assert attendus["total_net"] == Decimal("2850.00")
    assert attendus["total_net"] == Decimal(facture.total_net)


def test_avance_reportee(db, jeu, facture_emise):
    ws = ouvrir(db, facture_emise)[0].active
    assert ws["G32"].value == 250.0


def test_debordement_labor_et_odc(db, jeu):
    """Au-delà de 8 LABOR et 3 ODC : lignes insérées, formules réécrites."""
    nb_labor, nb_odc = CAPACITE_LABOR + 3, CAPACITE_ODC + 2
    with db.get_session() as session:
        for index in range(nb_labor):
            creer_entree(
                session, jeu["expert"], jeu["expert"], jeu["projet"],
                date(2026, 10, index + 1), Unite.HEURE, Decimal("1"),
                TypeActivite.FEDERAL, note_livrable=f"n{index}",
            )
        ids = [e.id for e in entrees_facturables(session, jeu["expert"])]
        facture, erreurs = creer_brouillon(session, jeu["expert"], jeu["expert"], entry_ids=ids)
        assert not erreurs, erreurs
        invoice_id = facture.id
        for index in range(nb_odc):
            ajouter_odc(
                session, jeu["expert"], invoice_id, "P", f"odc{index}", 1, Decimal("10")
            )
    with db.get_session() as session:
        facture, erreurs = emettre(session, jeu["expert"], invoice_id, date(2026, 10, 31))
        assert not erreurs, erreurs

    ws = ouvrir(db, invoice_id)[0].active
    k, j = nb_labor - CAPACITE_LABOR, nb_odc - CAPACITE_ODC
    labor_sous_total = 24 + k
    odc_premiere, odc_derniere = 27 + k, 29 + k + j
    odc_sous_total = 30 + k + j

    assert [ws.cell(row=LABOR_PREMIERE + i, column=2).value for i in range(nb_labor)] == list(
        range(1, nb_labor + 1)
    )
    assert ws.cell(row=odc_premiere, column=7).value == f"=E{odc_premiere}*F{odc_premiere}"
    assert ws.cell(row=odc_derniere, column=7).value == f"=E{odc_derniere}*F{odc_derniere}"
    assert ws.cell(row=odc_sous_total, column=7).value == (
        f"=SUM(G{odc_premiere}:G{odc_derniere})"
    )
    assert ws.cell(row=odc_sous_total + 1, column=7).value == (
        f"=G{labor_sous_total}+G{odc_sous_total}"
    )
    assert ws.cell(row=odc_sous_total + 3, column=7).value == (
        f"=G{odc_sous_total + 1}-G{odc_sous_total + 2}"
    )
    # Table étendue, signatures et mentions décalées sans altération.
    assert dict(ws.tables.items())["InvoiceDetails"] == f"B15:G{23 + k}"
    assert ws.cell(row=34 + k + j, column=2).value.startswith("Date:")
    assert ws.cell(row=35 + k + j, column=2).value.startswith("Please pdf this form")
    assert str(list(ws.merged_cells.ranges)[0]) in (
        f"A{37 + k + j}:H{37 + k + j}", "B2:E6",
    )
    assert f"A{37 + k + j}:H{37 + k + j}" in {str(m) for m in ws.merged_cells.ranges}
    assert ws.cell(row=37 + k + j, column=1).value == "Thank you for your business!"


def test_le_modele_d_origine_n_est_jamais_modifie(db, jeu, facture_emise):
    chemin = TEMPLATES_DIR / MODELE
    avant = chemin.read_bytes()
    ouvrir(db, facture_emise)
    assert chemin.read_bytes() == avant


def test_facture_inconnue(db):
    with db.get_session() as session:
        with pytest.raises(ValueError):
            generer_facture(session, 9999)
