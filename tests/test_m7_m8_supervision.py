"""M7 et M8 — métriques du dashboard, seuils, alertes et demandes d'extras."""
from __future__ import annotations

from datetime import date
from decimal import Decimal

import openpyxl

from core.alerts_service import (
    STATUT_ALERTE,
    STATUT_CRITIQUE,
    STATUT_EPUISE,
    STATUT_OK,
    brouillon_demande_extras,
    date_projetee_epuisement,
    depassement_a_signaler,
    exporter_vue_xlsx,
    generer_alertes,
    lister_alertes,
    marquer_alerte_traitee,
    statut_couple,
    tableau_supervision,
)
from core.config_service import definir_allocation
from core.entries_service import creer_entree
from core.enums import NiveauAlerte, StatutAlerte, TypeActivite, Unite
from core.integrity import calculer_consommation
from core.metrics_service import (
    avancement_par_projet,
    consommation_par_expert,
    facture_par_mois,
    kpis_factures,
    kpis_projets,
    kpis_taches,
    repartition_heures,
    timeline_livrables,
)


# --------------------------------------------------------------------------- #
# M7
# --------------------------------------------------------------------------- #
def test_kpis_taches_perimetre_expert(db, jeu, entrees):
    with db.get_session() as session:
        creer_entree(
            session, jeu["autre"], jeu["autre"], jeu["projet"], date(2026, 10, 9),
            Unite.HEURE, Decimal("8"), TypeActivite.FEDERAL,
        )
    with db.get_session() as session:
        siens = kpis_taches(session, user_id=jeu["expert"])
        tous = kpis_taches(session, user_id=None)
    assert siens["nb_entrees"] == 3
    assert siens["heures"] == Decimal("31.00")  # 7,5 + 16 + 7,5
    assert siens["jours"] == Decimal("3.88")
    assert siens["montant_facturable"] == Decimal("3100.00")
    assert tous["nb_entrees"] == 4
    assert tous["heures"] == Decimal("39.00")


def test_kpis_taches_concordent_avec_la_regle_d_integrite(db, jeu, entrees):
    """Les jours du dashboard sont ceux de la règle 5."""
    with db.get_session() as session:
        kpis = kpis_taches(session, user_id=jeu["expert"])
        couples = calculer_consommation(session, user_id=jeu["expert"])
    assert kpis["jours"] == sum(c.jours_consommes for c in couples).quantize(Decimal("0.01"))


def test_kpis_factures(db, jeu, facture_emise):
    from core.invoice_service import marquer_payee

    with db.get_session() as session:
        kpis = kpis_factures(session, user_id=jeu["expert"])
    assert kpis["par_statut"]["emise"]["nb"] == 1
    assert kpis["montant_facture"] == Decimal("2850.00")
    assert kpis["montant_en_attente"] == Decimal("2850.00")
    assert kpis["montant_paye"] == Decimal("0")

    with db.get_session() as session:
        marquer_payee(session, jeu["expert"], facture_emise)
    with db.get_session() as session:
        kpis = kpis_factures(session, user_id=jeu["expert"])
    assert kpis["montant_paye"] == Decimal("2850.00")
    assert kpis["montant_en_attente"] == Decimal("0")
    assert kpis["montant_facture"] == Decimal("2850.00")


def test_kpis_projets(db, jeu, entrees):
    with db.get_session() as session:
        kpis = kpis_projets(session)
    assert kpis["nb_projets_actifs"] == 2
    assert kpis["nb_livrables"] == 1
    assert kpis["nb_livrables_livres"] == 0


def test_avancement_par_projet(db, jeu, entrees):
    with db.get_session() as session:
        lignes = {l["projet"]: l for l in avancement_par_projet(session, user_id=jeu["expert"])}
    etude = lignes["Education Reform Study"]
    assert etude["budget_jours"] == Decimal("50.00")
    assert etude["jours_consommes"] == Decimal("3.88")
    assert etude["pct"] == Decimal("7.76")


def test_facture_par_mois(db, jeu, facture_emise):
    with db.get_session() as session:
        lignes = facture_par_mois(session, user_id=jeu["expert"])
    assert lignes == [
        {
            "annee": 2026, "mois": 10, "periode": "2026-10",
            "montant_emis": Decimal("2850.00"), "montant_paye": Decimal("0.00"),
        }
    ]


def test_repartition_heures(db, jeu, entrees):
    with db.get_session() as session:
        creer_entree(
            session, jeu["expert"], jeu["expert"], jeu["projet_b"], date(2026, 10, 10),
            Unite.HEURE, Decimal("4"), TypeActivite.BD,
        )
    with db.get_session() as session:
        par_activite = {l["libelle"]: l["heures"] for l in repartition_heures(
            session, "type_activite", user_id=jeu["expert"]
        )}
        par_projet = {l["libelle"]: l["heures"] for l in repartition_heures(
            session, "projet", user_id=jeu["expert"]
        )}
    assert par_activite == {"federal": Decimal("31.00"), "bd": Decimal("4.00")}
    assert par_projet["Education Reform Study"] == Decimal("31.00")
    assert par_projet["Proposal: Literacy Baseline"] == Decimal("4.00")


def test_timeline_livrables(db, jeu, entrees):
    with db.get_session() as session:
        lignes = timeline_livrables(session, user_id=jeu["expert"])
    assert len(lignes) == 1
    assert lignes[0]["intitule"] == "Inception report"
    assert lignes[0]["echeance"] == date(2026, 11, 30)


def test_consommation_par_expert(db, jeu, entrees):
    with db.get_session() as session:
        lignes = {l["expert"]: l for l in consommation_par_expert(session)}
    jane = lignes["Jane Doe"]
    assert jane["jours_alloues"] == Decimal("24.00")  # 20 + 4
    assert jane["jours_consommes"] == Decimal("3.88")


# --------------------------------------------------------------------------- #
# M8
# --------------------------------------------------------------------------- #
def test_statut_selon_les_seuils():
    assert statut_couple(Decimal("50"), Decimal("10"), 80, 95) == STATUT_OK
    assert statut_couple(Decimal("80"), Decimal("4"), 80, 95) == STATUT_ALERTE
    assert statut_couple(Decimal("95"), Decimal("1"), 80, 95) == STATUT_CRITIQUE
    assert statut_couple(Decimal("100"), Decimal("0"), 80, 95) == STATUT_EPUISE
    # L'épuisement prime sur le pourcentage.
    assert statut_couple(Decimal("10"), Decimal("-1"), 80, 95) == STATUT_EPUISE


def test_franchissement_de_seuil_fait_basculer_le_statut(db, jeu):
    """4 jours alloués sur le projet B : 3,2 j = 80 %, 3,9 j = 97,5 %."""
    with db.get_session() as session:
        creer_entree(
            session, jeu["expert"], jeu["expert"], jeu["projet_b"], date(2026, 10, 1),
            Unite.JOUR, Decimal("3.2"), TypeActivite.BD,
        )
    with db.get_session() as session:
        ligne = next(
            l for l in tableau_supervision(session, project_id=jeu["projet_b"])
            if l.user_id == jeu["expert"]
        )
        assert ligne.pct == Decimal("80.00")
        assert ligne.statut == STATUT_ALERTE

    with db.get_session() as session:
        creer_entree(
            session, jeu["expert"], jeu["expert"], jeu["projet_b"], date(2026, 10, 2),
            Unite.JOUR, Decimal("0.7"), TypeActivite.BD,
        )
    with db.get_session() as session:
        ligne = next(
            l for l in tableau_supervision(session, project_id=jeu["projet_b"])
            if l.user_id == jeu["expert"]
        )
        assert ligne.pct == Decimal("97.50")
        assert ligne.statut == STATUT_CRITIQUE


def test_epuisement_et_signalement_de_depassement(db, jeu):
    with db.get_session() as session:
        creer_entree(
            session, jeu["expert"], jeu["expert"], jeu["projet_b"], date(2026, 10, 1),
            Unite.JOUR, Decimal("4"), TypeActivite.BD,
        )
    with db.get_session() as session:
        ligne = next(
            l for l in tableau_supervision(session, project_id=jeu["projet_b"])
            if l.user_id == jeu["expert"]
        )
        assert ligne.statut == STATUT_EPUISE
        assert ligne.jours_restants == Decimal("0.00")
        message = depassement_a_signaler(session, jeu["expert"], jeu["projet_b"])
        assert message and "épuisée" in message.lower()


def test_couple_sans_allocation_est_en_depassement(db, jeu):
    """Un projet non alloué : saisie impossible, mais une allocation à 0 l'est."""
    with db.get_session() as session:
        definir_allocation(session, jeu["admin"], jeu["autre"], jeu["projet_b"], Decimal("0"))
        creer_entree(
            session, jeu["autre"], jeu["autre"], jeu["projet_b"], date(2026, 10, 1),
            Unite.HEURE, Decimal("2"), TypeActivite.BD,
        )
    with db.get_session() as session:
        message = depassement_a_signaler(session, jeu["autre"], jeu["projet_b"])
        assert message and "aucune allocation" in message.lower()


def test_tri_par_criticite(db, jeu):
    with db.get_session() as session:
        creer_entree(
            session, jeu["expert"], jeu["expert"], jeu["projet_b"], date(2026, 10, 1),
            Unite.JOUR, Decimal("4"), TypeActivite.BD,
        )
        creer_entree(
            session, jeu["expert"], jeu["expert"], jeu["projet"], date(2026, 10, 1),
            Unite.JOUR, Decimal("1"), TypeActivite.FEDERAL,
        )
    with db.get_session() as session:
        lignes = tableau_supervision(session, user_id=jeu["expert"])
    assert lignes[0].statut == STATUT_EPUISE
    assert lignes[-1].statut == STATUT_OK


def test_generation_d_alertes_idempotente(db, jeu):
    with db.get_session() as session:
        creer_entree(
            session, jeu["expert"], jeu["expert"], jeu["projet_b"], date(2026, 10, 1),
            Unite.JOUR, Decimal("4"), TypeActivite.BD,
        )
    with db.get_session() as session:
        premieres = generer_alertes(session, jeu["deputy"])
        assert len(premieres) == 1
        assert NiveauAlerte(premieres[0].niveau) is NiveauAlerte.EPUISE
    with db.get_session() as session:
        assert generer_alertes(session, jeu["deputy"]) == []
        assert len(lister_alertes(session, StatutAlerte.OUVERTE)) == 1


def test_alerte_marquee_traitee(db, jeu):
    with db.get_session() as session:
        creer_entree(
            session, jeu["expert"], jeu["expert"], jeu["projet_b"], date(2026, 10, 1),
            Unite.JOUR, Decimal("4"), TypeActivite.BD,
        )
    with db.get_session() as session:
        alerte = generer_alertes(session, jeu["deputy"])[0]
        alert_id = alerte.id
    with db.get_session() as session:
        assert marquer_alerte_traitee(session, jeu["deputy"], alert_id) == []
    with db.get_session() as session:
        assert lister_alertes(session, StatutAlerte.OUVERTE) == []
        assert len(lister_alertes(session, StatutAlerte.TRAITEE)) == 1


def test_projection_d_epuisement(db, jeu):
    """Rythme récent non nul et allocation restante : une date est proposée."""
    with db.get_session() as session:
        creer_entree(
            session, jeu["expert"], jeu["expert"], jeu["projet"], date.today(),
            Unite.JOUR, Decimal("5"), TypeActivite.FEDERAL,
        )
    with db.get_session() as session:
        projetee = date_projetee_epuisement(session, jeu["expert"], jeu["projet"])
        assert projetee is not None
        assert projetee > date.today()
        # Aucune saisie récente sur le projet B : pas de projection.
        assert date_projetee_epuisement(session, jeu["expert"], jeu["projet_b"]) is None


def test_demande_d_extras_est_un_brouillon(db, jeu, entrees):
    with db.get_session() as session:
        texte, erreurs = brouillon_demande_extras(
            session, jeu["deputy"], jeu["expert"], jeu["projet"], Decimal("5"),
            "Charge supplémentaire sur le volet quantitatif.",
        )
    assert not erreurs
    assert "Education Reform Study" in texte
    assert "Jane Doe" in texte
    assert "Jours supplémentaires demandés : 5" in texte
    assert "n'a pas été envoyé" in texte


def test_demande_d_extras_valide_ses_entrees(db, jeu):
    with db.get_session() as session:
        _, erreurs = brouillon_demande_extras(
            session, jeu["deputy"], jeu["expert"], jeu["projet"], Decimal("0"), "x"
        )
        assert erreurs
        _, erreurs = brouillon_demande_extras(
            session, jeu["deputy"], jeu["expert"], jeu["projet"], Decimal("5"), "  "
        )
        assert erreurs
        _, erreurs = brouillon_demande_extras(
            session, jeu["deputy"], 9999, jeu["projet"], Decimal("5"), "x"
        )
        assert erreurs == ["Expert inconnu."]


def test_export_de_la_vue(db, jeu, entrees):
    with db.get_session() as session:
        lignes = tableau_supervision(session)
        flux = exporter_vue_xlsx(lignes)
    ws = openpyxl.load_workbook(flux).active
    assert ws.title == "Supervision"
    assert [c.value for c in ws[1]] == [
        "Expert", "Projet", "Jours alloués", "Jours consommés",
        "Jours restants", "%", "Statut",
    ]
    assert ws.max_row == len(lignes) + 1


def test_supervision_sur_une_periode(db, jeu, entrees):
    """Borner la période restreint la consommation, pas l'allocation."""
    with db.get_session() as session:
        complete = next(
            l for l in tableau_supervision(session, user_id=jeu["expert"])
            if l.project_id == jeu["projet"]
        )
        bornee = next(
            l for l in tableau_supervision(
                session, user_id=jeu["expert"], debut=date(2026, 10, 2), fin=date(2026, 10, 31)
            )
            if l.project_id == jeu["projet"]
        )
    assert complete.jours_consommes == Decimal("3.88")
    assert bornee.jours_consommes == Decimal("2.94")  # 2 j + 7,5 h / 8
    assert bornee.jours_alloues == complete.jours_alloues
