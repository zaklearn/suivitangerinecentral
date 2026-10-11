"""M4 — critères d'acceptation : fidélité du modèle, formule, débordement."""
from __future__ import annotations

from datetime import date
from decimal import Decimal

import openpyxl
import pytest

from core.entries_service import creer_entree
from core.enums import TypeActivite, Unite
from core.timesheet_export import (
    CAPACITE_MODELE,
    LIGNE_PREMIERE_DONNEE,
    LIGNE_TOTAL,
    MODELE,
    generer_timesheet,
    libelle_mois,
    nom_fichier_timesheet,
    total_heures_attendu,
)
from core.xlsx_utils import TEMPLATES_DIR, charger_modele

COLONNES = range(1, 12)


def signature_style(cellule) -> tuple:
    police, bordure, alignement, fond = (
        cellule.font, cellule.border, cellule.alignment, cellule.fill,
    )
    return (
        police.name, police.sz, police.b, police.i,
        fond.patternType,
        bordure.left.style, bordure.right.style, bordure.top.style, bordure.bottom.style,
        alignement.horizontal, alignement.vertical, alignement.wrap_text,
    )


def ouvrir_genere(db, user_id, annee=2026, mois=10):
    with db.get_session() as session:
        flux, nom = generer_timesheet(session, user_id, annee, mois)
        attendu = total_heures_attendu(session, user_id, annee, mois)
    return openpyxl.load_workbook(flux), nom, attendu


def test_libelle_mois_independant_de_la_locale():
    assert libelle_mois(2026, 10) == "October 2026"
    assert libelle_mois(2026, 1) == "January 2026"
    with pytest.raises(ValueError):
        libelle_mois(2026, 13)


def test_nom_de_fichier():
    assert nom_fichier_timesheet("Jane Doe", 2026, 10) == "Timesheet_Jane_Doe_2026-10.xlsx"
    assert nom_fichier_timesheet("Élodie / Marc", 2026, 3).endswith("_2026-03.xlsx")


def test_entetes_et_identite(db, jeu, entrees):
    ws = ouvrir_genere(db, jeu["expert"])[0].active
    assert ws.title == "October 2026"
    assert ws["B3"].value == "Jane Doe"
    assert ws["B4"].value == "October 2026"
    assert ws["A1"].value == "Tangerine Central LLC - Time and Effort Record"
    assert ws["A8"].value == "Date"
    assert ws["D8"].value == "Number of Hours"
    assert ws["E8"].value == "Notes/Activities/Deliverables"
    assert str(list(ws.merged_cells.ranges)[0]) == "A6:E7"


def test_lignes_de_donnees(db, jeu, entrees):
    ws = ouvrir_genere(db, jeu["expert"])[0].active
    lignes = [
        [ws.cell(row=r, column=c).value for c in range(1, 6)]
        for r in range(LIGNE_PREMIERE_DONNEE, LIGNE_PREMIERE_DONNEE + 3)
    ]
    assert lignes[0][1] == "Education Reform Study"
    assert lignes[0][2] == "TC-2026-01"
    assert lignes[0][3] == 7.5
    assert lignes[0][4] == "Draft chapter 1"
    # L'entrée en jours est convertie par son snapshot : 2 j x 8 h.
    assert lignes[1][3] == 16.0
    assert ws.cell(row=LIGNE_PREMIERE_DONNEE, column=1).number_format == "mm/dd/yyyy"


def test_total_reste_une_formule_coherente(db, jeu, entrees):
    classeur, _, attendu = ouvrir_genere(db, jeu["expert"])
    ws = classeur.active
    assert ws.cell(row=LIGNE_TOTAL, column=4).value == "=SUBTOTAL(109,D9:D39)"
    assert ws.cell(row=LIGNE_TOTAL, column=3).value == "TOTAL HOURS:"
    assert attendu == Decimal("31.00")
    somme = sum(
        ws.cell(row=r, column=4).value or 0
        for r in range(LIGNE_PREMIERE_DONNEE, LIGNE_TOTAL)
    )
    assert Decimal(str(somme)) == attendu


def test_conversion_par_le_snapshot_et_non_la_config_courante(db, jeu, entrees):
    """Changer heures_par_jour après coup ne change pas le timesheet."""
    from core.config_service import maj_parametres

    with db.get_session() as session:
        maj_parametres(
            session, jeu["admin"], heures_par_jour=Decimal("6"),
            seuil_alerte_pct=80, seuil_critique_pct=95, prefixe_facture="TC",
        )
    ws = ouvrir_genere(db, jeu["expert"])[0].active
    # 2 jours restent convertis à 8 h/jour, le snapshot de l'entrée.
    assert ws.cell(row=LIGNE_PREMIERE_DONNEE + 1, column=4).value == 16.0


def test_toute_l_activite_est_reportee(db, jeu, entrees):
    """Une entrée non facturable, de type admin, figure quand même."""
    with db.get_session() as session:
        creer_entree(
            session, jeu["expert"], jeu["expert"], jeu["projet"], date(2026, 10, 15),
            Unite.HEURE, Decimal("3"), TypeActivite.ADMIN,
            note_livrable="Internal admin", facturable=False,
        )
    classeur, _, attendu = ouvrir_genere(db, jeu["expert"])
    ws = classeur.active
    notes = [
        ws.cell(row=r, column=5).value
        for r in range(LIGNE_PREMIERE_DONNEE, LIGNE_PREMIERE_DONNEE + 4)
    ]
    assert "Internal admin" in notes
    assert attendu == Decimal("34.00")


def test_debordement_au_dela_de_31_entrees(db, jeu):
    """Lignes insérées, formule étendue, signatures préservées."""
    nb = CAPACITE_MODELE + 4
    with db.get_session() as session:
        for index in range(nb):
            _, erreurs = creer_entree(
                session, jeu["expert"], jeu["expert"], jeu["projet"],
                date(2026, 10, (index % 28) + 1), Unite.HEURE, Decimal("1"),
                TypeActivite.FEDERAL, note_livrable=f"n{index}",
            )
            assert not erreurs, erreurs

    classeur, _, attendu = ouvrir_genere(db, jeu["expert"])
    ws = classeur.active
    decalage = nb - CAPACITE_MODELE
    derniere = 39 + decalage
    total = LIGNE_TOTAL + decalage

    assert ws.cell(row=total, column=4).value == f"=SUBTOTAL(109,D9:D{derniere})"
    assert ws.cell(row=total, column=3).value == "TOTAL HOURS:"
    assert attendu == Decimal(nb)
    # Toutes les lignes de données portent une valeur, aucune n'est vide.
    assert all(
        ws.cell(row=r, column=4).value is not None
        for r in range(LIGNE_PREMIERE_DONNEE, derniere + 1)
    )
    # Certification et signatures suivent le décalage, intactes.
    assert ws.cell(row=41 + decalage, column=2).value.startswith("I certify that this record")
    assert ws.cell(row=43 + decalage, column=4).value.startswith("Signature:")
    assert ws.cell(row=47 + decalage, column=2).value.startswith("Date:")
    # La table Excel est étendue d'autant.
    assert dict(ws.tables.items())["Table224567893"] == f"A8:E{48 + decalage}"


def test_fidelite_de_style(db, jeu):
    """Hors colonne Date, aucun style ne diffère du modèle — débordement inclus."""
    nb = CAPACITE_MODELE + 4
    with db.get_session() as session:
        for index in range(nb):
            creer_entree(
                session, jeu["expert"], jeu["expert"], jeu["projet"],
                date(2026, 10, (index % 28) + 1), Unite.HEURE, Decimal("1"),
                TypeActivite.FEDERAL,
            )
    decalage = nb - CAPACITE_MODELE
    modele = charger_modele(MODELE).active
    genere = ouvrir_genere(db, jeu["expert"])[0].active

    ecarts = []
    # En-tête et instructions : position inchangée.
    for r in range(1, LIGNE_PREMIERE_DONNEE):
        for c in COLONNES:
            if signature_style(modele.cell(row=r, column=c)) != signature_style(
                genere.cell(row=r, column=c)
            ):
                ecarts.append(f"haut L{r}C{c}")
    # Lignes insérées : habillage de la dernière ligne de données du modèle.
    for r in range(40, 40 + decalage):
        for c in COLONNES:
            if signature_style(modele.cell(row=39, column=c)) != signature_style(
                genere.cell(row=r, column=c)
            ):
                ecarts.append(f"insérée L{r}C{c}")
    # Total, certification, signatures : décalés mais identiques.
    for r in range(LIGNE_TOTAL, modele.max_row + 1):
        for c in COLONNES:
            if signature_style(modele.cell(row=r, column=c)) != signature_style(
                genere.cell(row=r + decalage, column=c)
            ):
                ecarts.append(f"bas L{r}->L{r + decalage}C{c}")
    assert ecarts == []

    for lettre, dimension in modele.column_dimensions.items():
        assert genere.column_dimensions[lettre].width == dimension.width


def test_le_modele_d_origine_n_est_jamais_modifie(db, jeu, entrees):
    chemin = TEMPLATES_DIR / MODELE
    avant = chemin.read_bytes()
    ouvrir_genere(db, jeu["expert"])
    assert chemin.read_bytes() == avant


def test_expert_inconnu(db):
    with db.get_session() as session:
        with pytest.raises(ValueError):
            generer_timesheet(session, 9999, 2026, 10)
